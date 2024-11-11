import speech_recognition as sr
from pydub import AudioSegment
from docx import Document

audio = AudioSegment.from_mp3("C:/Users/19990772/Desktop/Codigos/Transcripcion/pythonProject1/audio.mp3")

"""
audio.export("Estudiantes 2022-2024 Nutrición y dietética.wav", format="wav")

recognizer = sr.Recognizer()

# Cargar el archivo WAV
with sr.AudioFile("Estudiantes 2022-2024 Nutrición y dietética.wav") as source:
    audio_data = recognizer.record(source)

# Reconocer el audio
try:
    texto = recognizer.recognize_google(audio_data, language='es-ES')
    print("Texto transcrito:", texto)

    # Guardar en un archivo Word
    doc = Document()
    doc.add_heading('Transcripción de audio', level=1)
    doc.add_paragraph(texto)
    doc.save("transcripcion.docx")

except sr.UnknownValueError:
    print("No se pudo entender el audio")
except sr.RequestError as e:
    print(f"No se pudo solicitar resultados; {e}")
"""